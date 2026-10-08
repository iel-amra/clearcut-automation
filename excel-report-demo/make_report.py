#!/usr/bin/env python3
"""Turn a raw sales CSV into a formatted Excel report with pivot-style
summaries, charts and conditional formatting.

Usage:
    python make_report.py sales.csv report.xlsx
"""
import sys

import pandas as pd
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
HEADER_FONT = Font(bold=True, color="FFFFFF")
MONEY = '#,##0.00 €'


def style_header(ws, ncols: int) -> None:
    for c in range(1, ncols + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "A2"


def autosize(ws) -> None:
    for col in ws.columns:
        width = max(len(str(c.value)) if c.value is not None else 0 for c in col) + 2
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(width, 10), 50)


def build(csv_path: str, out_path: str) -> None:
    df = pd.read_csv(csv_path, parse_dates=["date"])
    df["revenue"] = df["quantity"] * df["unit_price"]
    df["month"] = df["date"].dt.to_period("M").astype(str)

    by_month = df.groupby("month", as_index=False)["revenue"].sum()
    by_product = (df.groupby("product", as_index=False)
                    .agg(units=("quantity", "sum"), revenue=("revenue", "sum"))
                    .sort_values("revenue", ascending=False))
    by_region = df.pivot_table(index="region", columns="month", values="revenue", aggfunc="sum", fill_value=0).reset_index()

    with pd.ExcelWriter(out_path, engine="openpyxl") as xw:
        summary = pd.DataFrame({
            "metric": ["Total revenue", "Orders", "Average order value", "Best product", "Best region"],
            "value": [round(df["revenue"].sum(), 2), len(df), round(df["revenue"].mean(), 2),
                      by_product.iloc[0]["product"], df.groupby("region")["revenue"].sum().idxmax()],
        })
        summary.to_excel(xw, sheet_name="Summary", index=False)
        by_month.to_excel(xw, sheet_name="By month", index=False)
        by_product.to_excel(xw, sheet_name="By product", index=False)
        by_region.to_excel(xw, sheet_name="Region x month", index=False)
        df.drop(columns="month").to_excel(xw, sheet_name="Raw data", index=False)

        wb = xw.book
        for name in wb.sheetnames:
            ws = wb[name]
            style_header(ws, ws.max_column)
            autosize(ws)
            ws.auto_filter.ref = ws.dimensions

        ws = wb["Summary"]
        for r in (2, 4):
            ws.cell(row=r, column=2).number_format = MONEY
        ws.column_dimensions["B"].width = 22

        ws = wb["By month"]
        for r in range(2, ws.max_row + 1):
            ws.cell(row=r, column=2).number_format = MONEY
        chart = LineChart()
        chart.title, chart.y_axis.title, chart.x_axis.title = "Revenue per month", "€", "Month"
        chart.add_data(Reference(ws, min_col=2, min_row=1, max_row=ws.max_row), titles_from_data=True)
        chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=ws.max_row))
        chart.height, chart.width = 8, 18
        ws.add_chart(chart, "D2")

        ws = wb["By product"]
        for r in range(2, ws.max_row + 1):
            ws.cell(row=r, column=3).number_format = MONEY
        bar = BarChart()
        bar.type, bar.title, bar.y_axis.title = "bar", "Revenue per product", "€"
        bar.add_data(Reference(ws, min_col=3, min_row=1, max_row=ws.max_row), titles_from_data=True)
        bar.set_categories(Reference(ws, min_col=1, min_row=2, max_row=ws.max_row))
        bar.height, bar.width = 8, 18
        ws.add_chart(bar, "E2")

        ws = wb["Region x month"]
        rng = f"B2:{get_column_letter(ws.max_column)}{ws.max_row}"
        for row in ws[rng]:
            for cell in row:
                cell.number_format = MONEY
        ws.conditional_formatting.add(rng, ColorScaleRule(start_type="min", start_color="F8696B",
                                                           mid_type="percentile", mid_value=50, mid_color="FFEB84",
                                                           end_type="max", end_color="63BE7B"))
    print(f"report written: {out_path}")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    build(sys.argv[1], sys.argv[2])
